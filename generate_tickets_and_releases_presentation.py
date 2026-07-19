#!/usr/bin/env python3
"""Generate pilot defects presentation from RTR CSV data.

Style and layout: ai/rules/cherkizovo-presentations.md
Theme: themes/cherkizovo.py
"""

from __future__ import annotations

import argparse
import csv
import math
import sys
from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, MSO_UNDERLINE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

ROOT_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT_DIR))

from themes.cherkizovo import (  # noqa: E402
    CHARS_PER_INCH,
    COLORS,
    COLUMN_WIDTHS_IN,
    CSV_EXCLUDE_TEXT,
    CSV_INCLUDE_TEXT,
    HEADER_ROW_HEIGHT_IN,
    LINE_HEIGHT_IN,
    LOGO_LEFT_IN,
    LOGO_TOP_IN,
    LOGO_WIDTH_IN,
    MAX_DATA_HEIGHT_IN,
    MIN_DATA_ROW_HEIGHT_IN,
    PAGE_FONT,
    PAGE_NUMBER_SIZE,
    SLIDE_HEIGHT_IN,
    SLIDE_TITLE_TEMPLATE,
    SLIDE_WIDTH_IN,
    STATUS_COL_WIDTH_IN,
    TABLE_BODY_SIZE,
    TABLE_COLUMNS,
    TABLE_FONT,
    TABLE_HEADER_SIZE,
    TABLE_LEFT_IN,
    TABLE_TOP_IN,
    TABLE_WIDTH_IN,
    TASK_COL_WIDTH_IN,
    TITLE_FONT,
    TITLE_LEFT_IN,
    TITLE_SIZE,
    TITLE_TOP_IN,
    TRACKER_BASE_URL,
    VALUE_COL_WIDTH_IN,
    apply_cherkizovo_theme,
)

ASSETS_DIR = Path(__file__).resolve().parent / "assets"
LOGO_PATH = ASSETS_DIR / "logo_0.png"


def rgb(name: str) -> RGBColor:
    red, green, blue = COLORS[name]
    return RGBColor(red, green, blue)


@dataclass(frozen=True)
class CherkizovoTheme:
    """Runtime theme mapped from cherkizovo-presentations.md."""

    title_font = TITLE_FONT
    table_font = TABLE_FONT
    page_font = PAGE_FONT

    title_rgb = rgb("title")
    text_rgb = rgb("black")
    key_rgb = rgb("key")
    accent_rgb = rgb("critical")
    header_bg_rgb = rgb("dark_red")
    header_fg_rgb = rgb("white")
    alt_row_rgb = rgb("alt_row")
    white_rgb = rgb("white")
    page_number_rgb = rgb("page_number")

    title_size = TITLE_SIZE
    header_size = TABLE_HEADER_SIZE
    body_size = TABLE_BODY_SIZE
    page_number_size = PAGE_NUMBER_SIZE


THEME = CherkizovoTheme()

SLIDE_WIDTH = Inches(SLIDE_WIDTH_IN)
SLIDE_HEIGHT = Inches(SLIDE_HEIGHT_IN)
TITLE_LEFT = Inches(TITLE_LEFT_IN)
TITLE_TOP = Inches(TITLE_TOP_IN)
TITLE_WIDTH = Inches(8.5)
LOGO_LEFT = Inches(LOGO_LEFT_IN)
LOGO_TOP = Inches(LOGO_TOP_IN)
LOGO_WIDTH = Inches(LOGO_WIDTH_IN)
TABLE_LEFT = Inches(TABLE_LEFT_IN)
TABLE_TOP = Inches(TABLE_TOP_IN)
TABLE_WIDTH = Inches(TABLE_WIDTH_IN)
HEADER_ROW_HEIGHT = Inches(HEADER_ROW_HEIGHT_IN)
CELL_MARGIN_LR = Pt(3)
CELL_MARGIN_TB = Pt(2)

COLUMNS = TABLE_COLUMNS
COLUMN_WIDTHS = [Inches(width) for width in COLUMN_WIDTHS_IN]


def load_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig") as handle:
        rows = list(csv.reader(handle))

    header = rows[5]
    index = {name.strip(): position for position, name in enumerate(header)}

    def get(row: list[str], name: str, default: str = "") -> str:
        position = index.get(name)
        if position is None or position >= len(row):
            return default
        return (row[position] or "").strip()

    records: list[dict[str, str]] = []
    for row in rows[6:]:
        team = get(row, "Команда")
        if not team:
            continue

        row_text = " | ".join(row).lower()
        if CSV_INCLUDE_TEXT not in row_text or CSV_EXCLUDE_TEXT in row_text:
            continue

        records.append(
            {
                "team": team,
                "defect": get(row, "Дефект"),
                "task_key": get(row, "Ключ задач в ЯТ (ссылка)"),
                "description": get(row, "Описание задачи (текстовое поле)"),
                "priority": get(row, "Приоритет"),
                "release": get(row, "Релиз"),
                "chtz_status": get(row, "Статус ЧТЗ"),
                "development": get(row, "Разработка"),
                "comments": get(row, "Корректировки/комментарии") or get(row, "Комментарии"),
                "cost": get(row, "Стоимость разработки, тыс.руб."),
                "value": get(row, "value"),
                "ppk_critical": get(row, "критичны к запуску ППК"),
            }
        )
    return records


def is_critical(record: dict[str, str]) -> bool:
    priority = record["priority"].lower()
    ppk = record["ppk_critical"].lower()
    comments = record["comments"].lower()
    return (
        "критич" in priority
        or "критич" in ppk
        or "критич" in comments
        or "важно" in comments
    )


def value_text(record: dict[str, str]) -> str:
    if record["value"]:
        return record["value"]
    comments = record["comments"]
    if comments and len(comments) > 35 and not comments.lower().startswith("ок"):
        return comments
    return ""


def status_text(record: dict[str, str]) -> str:
    development = record["development"]
    release = record["release"]
    chtz_status = record["chtz_status"]
    comments = record["comments"]
    analyst = record.get("analyst", "")

    lowered = development.lower()
    if "беринг" in lowered:
        if any(word in lowered for word in ("разработ", "код", "dev")):
            return "в разработке у БерингПро"
        return "на оценке у БерингПро"
    if "базис" in lowered:
        return "в разработке у Базис"
    if "1с-перспектива" in lowered or "перспектив" in lowered:
        return "на оценке у 1С-Перспектива"
    if development:
        return development

    parts: list[str] = []
    if chtz_status:
        parts.append(chtz_status)
    if release:
        parts.append(f"{release} релиз")
    if comments and len(comments) <= 50 and "критич" not in comments.lower():
        parts.append(comments)
    if parts:
        return ", ".join(parts)
    return "в работе"


def task_parts(record: dict[str, str]) -> list[tuple[str, bool]]:
    defect = record["defect"] or "Задача"
    key = record["task_key"] or "—"
    description = record["description"]
    text = f"{key}: {defect} - {description}"
    parts: list[tuple[str, bool]] = [(text, False)]

    lowered = f"{description} {record['comments']}".lower()
    if "упп" in lowered or "реализован" in lowered:
        parts.append((" – реализовано в УПП", True))
    return parts


def cost_text(record: dict[str, str]) -> str:
    cost = record["cost"]
    if cost in {"", "0", "#N/A"}:
        return ""
    return cost


def wrapped_lines(text: str, column_width_in: float) -> int:
    if not text:
        return 1
    chars_per_line = max(1, int(column_width_in * CHARS_PER_INCH))
    return max(1, math.ceil(len(text) / chars_per_line))


def estimate_row_height(record: dict[str, str]) -> float:
    task = "".join(text for text, _ in task_parts(record))
    value = value_text(record)
    status = status_text(record)
    if is_critical(record):
        status = f"{status}\nКритичная!"

    lines = max(
        wrapped_lines(task, TASK_COL_WIDTH_IN),
        wrapped_lines(value, VALUE_COL_WIDTH_IN),
        wrapped_lines(status, STATUS_COL_WIDTH_IN),
    )
    return max(MIN_DATA_ROW_HEIGHT_IN, lines * LINE_HEIGHT_IN + 0.08)


def finalize_row_heights(natural_heights: list[float]) -> list[float]:
    if not natural_heights:
        return []
    natural_total = sum(natural_heights)
    if natural_total <= MAX_DATA_HEIGHT_IN:
        return natural_heights
    ratio = MAX_DATA_HEIGHT_IN / natural_total
    return [height * ratio for height in natural_heights]


def pack_slides(records: list[dict[str, str]]) -> list[tuple[list[dict[str, str]], list[float]]]:
    slides: list[tuple[list[dict[str, str]], list[float]]] = []
    current_records: list[dict[str, str]] = []
    current_heights: list[float] = []

    for record in records:
        row_height = estimate_row_height(record)
        used_height = sum(current_heights)
        if current_records and used_height + row_height > MAX_DATA_HEIGHT_IN:
            slides.append((current_records, finalize_row_heights(current_heights)))
            current_records = []
            current_heights = []

        current_records.append(record)
        current_heights.append(row_height)

    if current_records:
        slides.append((current_records, finalize_row_heights(current_heights)))

    return slides


def set_cell_border(cell, *, color: str = "000000", width: str = "6350") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    for edge in ("lnL", "lnR", "lnT", "lnB"):
        tag = qn(f"a:{edge}")
        element = tc_pr.find(tag)
        if element is None:
            element = OxmlElement(f"a:{edge}")
            tc_pr.append(element)
        element.set("w", width)
        element.set("cap", "flat")
        element.set("cmpd", "sng")
        element.set("algn", "ctr")
        fill = element.find(qn("a:solidFill"))
        if fill is None:
            fill = OxmlElement("a:solidFill")
            element.append(fill)
        srgb = fill.find(qn("a:srgbClr"))
        if srgb is None:
            srgb = OxmlElement("a:srgbClr")
            fill.append(srgb)
        srgb.set("val", color)


def style_table_borders(table) -> None:
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(cell)


def tracker_url(task_key: str) -> str:
    return f"{TRACKER_BASE_URL}{task_key}"


def set_run_font(
    run,
    *,
    font_name: str,
    size: int,
    bold: bool = False,
    color: RGBColor | None = None,
) -> None:
    run.font.name = font_name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_textbox(
    text_frame,
    text: str,
    *,
    font_name: str,
    size: int,
    bold: bool = False,
    color: RGBColor | None = None,
    align: PP_ALIGN = PP_ALIGN.LEFT,
) -> None:
    text_frame.clear()
    paragraph = text_frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    set_run_font(
        run,
        font_name=font_name,
        size=size,
        bold=bold,
        color=color or THEME.text_rgb,
    )


def set_cell_runs(
    cell,
    runs: list[tuple[str, bool, RGBColor | None]],
    *,
    size: int = THEME.body_size,
    align=PP_ALIGN.LEFT,
    fill: RGBColor | None = None,
) -> None:
    cell.text = ""
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = CELL_MARGIN_LR
    cell.margin_right = CELL_MARGIN_LR
    cell.margin_top = CELL_MARGIN_TB
    cell.margin_bottom = CELL_MARGIN_TB
    cell.text_frame.word_wrap = True
    paragraph = cell.text_frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_before = Pt(0)
    paragraph.space_after = Pt(0)
    paragraph.line_spacing = 1.0
    for text, bold, color in runs:
        if not text:
            continue
        run = paragraph.add_run()
        run.text = text
        set_run_font(
            run,
            font_name=THEME.table_font,
            size=size,
            bold=bold,
            color=color or THEME.text_rgb,
        )
    if fill is not None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill


def set_key_cell(cell, task_key: str, *, fill: RGBColor) -> None:
    """Set the task-key cell with an optional Yandex Tracker hyperlink."""
    cell.text = ""
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = CELL_MARGIN_LR
    cell.margin_right = CELL_MARGIN_LR
    cell.margin_top = CELL_MARGIN_TB
    cell.margin_bottom = CELL_MARGIN_TB
    cell.text_frame.word_wrap = True
    paragraph = cell.text_frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.LEFT
    paragraph.space_before = Pt(0)
    paragraph.space_after = Pt(0)
    paragraph.line_spacing = 1.0

    run = paragraph.add_run()
    run.text = task_key
    set_run_font(
        run,
        font_name=THEME.table_font,
        size=THEME.body_size,
        color=THEME.key_rgb,
    )
    if task_key != "—":
        run.hyperlink.address = tracker_url(task_key)
        run.font.underline = MSO_UNDERLINE.SINGLE_LINE
        run.font.color.rgb = THEME.key_rgb

    cell.fill.solid()
    cell.fill.fore_color.rgb = fill


def set_cell(
    cell,
    text: str,
    *,
    font_name: str = THEME.table_font,
    bold: bool = False,
    size: int = THEME.body_size,
    align=PP_ALIGN.LEFT,
    fill: RGBColor | None = None,
    font_color: RGBColor | None = None,
) -> None:
    set_cell_runs(
        cell,
        [(str(text), bold, font_color)],
        size=size,
        align=align,
        fill=fill,
    )
    for run in cell.text_frame.paragraphs[0].runs:
        run.font.name = font_name


def add_logo(slide) -> None:
    if not LOGO_PATH.exists():
        return
    slide.shapes.add_picture(str(LOGO_PATH), LOGO_LEFT, LOGO_TOP, width=LOGO_WIDTH)


def add_slide_title(slide, page_index: int, total_pages: int) -> None:
    title_box = slide.shapes.add_textbox(TITLE_LEFT, TITLE_TOP, TITLE_WIDTH, Inches(0.45))
    set_textbox(
        title_box.text_frame,
        SLIDE_TITLE_TEMPLATE.format(page=page_index, total=total_pages),
        font_name=THEME.title_font,
        size=THEME.title_size,
        bold=True,
        color=THEME.title_rgb,
    )


def build_presentation(records: list[dict[str, str]], output_path: Path) -> None:
    slides = pack_slides(records)
    total_pages = max(1, len(slides))
    row_offset = 0

    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT
    apply_cherkizovo_theme(prs)
    blank = prs.slide_layouts[6]

    for page_index, (chunk, row_heights) in enumerate(slides, start=1):
        slide = prs.slides.add_slide(blank)
        add_logo(slide)
        add_slide_title(slide, page_index, total_pages)

        table_height = Inches(HEADER_ROW_HEIGHT_IN + sum(row_heights))

        table_shape = slide.shapes.add_table(
            len(chunk) + 1,
            len(COLUMNS),
            TABLE_LEFT,
            TABLE_TOP,
            TABLE_WIDTH,
            table_height,
        )
        table = table_shape.table
        for index, width in enumerate(COLUMN_WIDTHS):
            table.columns[index].width = width

        table.rows[0].height = HEADER_ROW_HEIGHT
        for row_index, row_height in enumerate(row_heights, start=1):
            table.rows[row_index].height = Inches(row_height)

        for column_index, title in enumerate(COLUMNS):
            set_cell(
                table.cell(0, column_index),
                title,
                font_name=THEME.table_font,
                bold=True,
                size=THEME.header_size,
                align=PP_ALIGN.CENTER,
                fill=THEME.header_bg_rgb,
                font_color=THEME.header_fg_rgb,
            )

        for row_index, record in enumerate(chunk, start=1):
            row_number = row_offset + row_index
            fill = THEME.alt_row_rgb if row_index % 2 == 0 else THEME.white_rgb
            critical = is_critical(record)
            accent = THEME.accent_rgb if critical else THEME.text_rgb

            set_cell(
                table.cell(row_index, 0),
                str(row_number),
                size=THEME.body_size,
                align=PP_ALIGN.CENTER,
                fill=fill,
            )
            task_key = record["task_key"] or "—"
            set_key_cell(table.cell(row_index, 1), task_key, fill=fill)

            task_runs = [(text, bold, THEME.text_rgb) for text, bold in task_parts(record)]
            set_cell_runs(table.cell(row_index, 2), task_runs, fill=fill)

            value = value_text(record)
            if len(value) > 320:
                value = value[:317] + "..."
            set_cell(table.cell(row_index, 3), value, fill=fill)

            set_cell(
                table.cell(row_index, 4),
                cost_text(record),
                bold=critical and bool(cost_text(record)),
                align=PP_ALIGN.RIGHT,
                fill=fill,
                font_color=accent if critical and cost_text(record) else THEME.text_rgb,
            )

            status = status_text(record)
            status_runs: list[tuple[str, bool, RGBColor | None]] = [(status, False, THEME.text_rgb)]
            if critical:
                status_runs.append(("\nКритичная!", True, THEME.accent_rgb))
            set_cell_runs(
                table.cell(row_index, 5),
                status_runs,
                align=PP_ALIGN.LEFT,
                fill=fill,
            )

        style_table_borders(table)
        row_offset += len(chunk)

    prs.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--csv",
        type=Path,
        default=Path("data/RTR-status.csv"),
        help="Path to source CSV file",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("Дефекты_и_разработки_по_Пилоту.pptx"),
        help="Path to generated presentation",
    )
    args = parser.parse_args()

    records = load_csv(args.csv)
    build_presentation(records, args.output)
    slides = pack_slides(records)
    print(f"Created {args.output} with {len(records)} tasks across {len(slides)} slides")


if __name__ == "__main__":
    main()
