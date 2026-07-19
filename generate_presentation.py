#!/usr/bin/env python3
"""Generate a Cherkizovo-styled backlog presentation from JSON data."""

from __future__ import annotations

import json
from collections import OrderedDict
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
DATA_PATH = ROOT / "data" / "tasks.json"
LOGO_PATH = ROOT / "assets" / "cherkizovo_logo.png"
OUTPUT_PATH = ROOT / "output" / "RTR-UU-backlog.pptx"

BRAND_RED = RGBColor(174, 23, 45)
BRAND_RED_LIGHT = RGBColor(230, 27, 64)
BLACK = RGBColor(0, 0, 0)
WHITE = RGBColor(255, 255, 255)
GRAY = RGBColor(80, 80, 80)

SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)

COLUMNS = [
    ("Категория", 1.15),
    ("Доработка", 0.65),
    ("Новая\nзадача", 0.65),
    ("Ключ / Описание", 4.6),
    ("Трудоемкость,\nч**", 1.0),
    ("Value", 3.9),
]


def load_data(path: Path) -> dict:
    with path.open(encoding="utf-8") as fh:
        return json.load(fh)


def group_tasks_by_category(tasks: list[dict]) -> OrderedDict[str, list[dict]]:
    grouped: OrderedDict[str, list[dict]] = OrderedDict()
    for task in tasks:
        grouped.setdefault(task["category"], []).append(task)
    return grouped


def set_cell_text(
    cell,
    text: str,
    *,
    bold: bool = False,
    size: int = 10,
    color: RGBColor = BLACK,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    valign: MSO_ANCHOR = MSO_ANCHOR.TOP,
) -> None:
    cell.text = ""
    paragraph = cell.text_frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    cell.text_frame.word_wrap = True
    cell.vertical_anchor = valign
    cell.margin_left = Pt(4)
    cell.margin_right = Pt(4)
    cell.margin_top = Pt(3)
    cell.margin_bottom = Pt(3)


def style_header_row(table) -> None:
    for col_idx, (title, _) in enumerate(COLUMNS):
        cell = table.cell(0, col_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = BRAND_RED
        set_cell_text(
            cell,
            title,
            bold=True,
            size=11,
            color=WHITE,
            align=PP_ALIGN.CENTER,
            valign=MSO_ANCHOR.MIDDLE,
        )


def add_value_text(cell, value: str) -> None:
    cell.text = ""
    paragraph = cell.text_frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.LEFT

    if value.startswith("ВАЖНО!"):
        important, rest = value.split("!", 1)
        run = paragraph.add_run()
        run.text = important + "!"
        run.font.name = "Calibri"
        run.font.size = Pt(10)
        run.font.bold = True
        run.font.color.rgb = BRAND_RED_LIGHT

        if rest.strip():
            run = paragraph.add_run()
            run.text = rest
            run.font.name = "Calibri"
            run.font.size = Pt(10)
            run.font.color.rgb = BLACK
    else:
        run = paragraph.add_run()
        run.text = value
        run.font.name = "Calibri"
        run.font.size = Pt(10)
        run.font.color.rgb = BLACK

    cell.text_frame.word_wrap = True
    cell.vertical_anchor = MSO_ANCHOR.TOP


def format_effort(task: dict) -> str:
    analytics = task.get("analytics_hours")
    dev = task.get("dev_hours")
    if analytics is not None and dev is not None:
        return f"{analytics}+{dev}"
    if analytics is not None:
        return str(analytics)
    if dev is not None:
        return str(dev)
    return ""


def format_description(task: dict) -> str:
    parts = []
    if task.get("task_key"):
        parts.append(task["task_key"])
    if task.get("dev_key"):
        parts.append(task["dev_key"])
    if task.get("description"):
        parts.append(task["description"])
    if task.get("complexity"):
        parts.append(f"Сложность: {task['complexity']}")
    return "\n".join(parts)


def add_slide_header(slide, title: str, page_number: int) -> None:
    if LOGO_PATH.exists():
        slide.shapes.add_picture(
            str(LOGO_PATH),
            left=Inches(11.55),
            top=Inches(0.18),
            width=Inches(1.55),
        )

    title_box = slide.shapes.add_textbox(
        Inches(0.45), Inches(0.25), Inches(10.5), Inches(0.55)
    )
    title_frame = title_box.text_frame
    title_frame.clear()
    paragraph = title_frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = title
    run.font.name = "Calibri"
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = BLACK

    page_box = slide.shapes.add_textbox(
        Inches(12.55), Inches(7.0), Inches(0.5), Inches(0.3)
    )
    page_frame = page_box.text_frame
    page_frame.clear()
    paragraph = page_frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.RIGHT
    run = paragraph.add_run()
    run.text = str(page_number)
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.font.color.rgb = BLACK


def add_footnotes(slide, footnotes: list[str]) -> None:
    foot_box = slide.shapes.add_textbox(
        Inches(0.45), Inches(6.95), Inches(8.5), Inches(0.45)
    )
    frame = foot_box.text_frame
    frame.clear()
    for idx, note in enumerate(footnotes):
        paragraph = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        run = paragraph.add_run()
        run.text = note
        run.font.name = "Calibri"
        run.font.size = Pt(9)
        run.font.color.rgb = GRAY


def build_table(slide, grouped_tasks: OrderedDict[str, list[dict]], top: float) -> None:
    total_rows = sum(len(items) for items in grouped_tasks.values())
    table = slide.shapes.add_table(
        total_rows + 1,
        len(COLUMNS),
        Inches(0.45),
        Inches(top),
        sum(width for _, width in COLUMNS),
        Inches(min(5.8, 0.42 * (total_rows + 1) + 0.5)),
    ).table

    col_offset = 0.0
    for col_idx, (_, width) in enumerate(COLUMNS):
        table.columns[col_idx].width = Inches(width)
        col_offset += width

    style_header_row(table)

    row_idx = 1
    for category, items in grouped_tasks.items():
        start_row = row_idx
        for item in items:
            set_cell_text(
                table.cell(row_idx, 0),
                category if row_idx == start_row else "",
                bold=True,
                size=10,
                valign=MSO_ANCHOR.MIDDLE,
            )
            set_cell_text(
                table.cell(row_idx, 1),
                "V" if item.get("is_refinement") else "",
                bold=True,
                size=12,
                align=PP_ALIGN.CENTER,
                valign=MSO_ANCHOR.MIDDLE,
            )
            set_cell_text(
                table.cell(row_idx, 2),
                "V" if item.get("is_new") else "",
                bold=True,
                size=12,
                align=PP_ALIGN.CENTER,
                valign=MSO_ANCHOR.MIDDLE,
            )
            set_cell_text(table.cell(row_idx, 3), format_description(item), size=9)
            set_cell_text(
                table.cell(row_idx, 4),
                format_effort(item),
                size=10,
                align=PP_ALIGN.CENTER,
                valign=MSO_ANCHOR.MIDDLE,
            )
            add_value_text(table.cell(row_idx, 5), item.get("value", ""))
            row_idx += 1

        if len(items) > 1:
            table.cell(start_row, 0).merge(table.cell(row_idx - 1, 0))


def chunk_grouped_tasks(
    grouped_tasks: OrderedDict[str, list[dict]], max_rows: int
) -> list[OrderedDict[str, list[dict]]]:
    chunks: list[OrderedDict[str, list[dict]]] = []
    current: OrderedDict[str, list[dict]] = OrderedDict()
    current_count = 0

    for category, items in grouped_tasks.items():
        if current_count + len(items) > max_rows and current:
            chunks.append(current)
            current = OrderedDict()
            current_count = 0
        current[category] = items
        current_count += len(items)

    if current:
        chunks.append(current)
    return chunks


def generate_presentation(data: dict, output_path: Path) -> Path:
    grouped = group_tasks_by_category(data["tasks"])
    chunks = chunk_grouped_tasks(grouped, max_rows=8)

    presentation = Presentation()
    presentation.slide_width = SLIDE_WIDTH
    presentation.slide_height = SLIDE_HEIGHT

    for page_idx, chunk in enumerate(chunks, start=1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        add_slide_header(slide, data["title"], page_idx)
        build_table(slide, chunk, top=1.0)
        add_footnotes(slide, data.get("footnotes", []))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output_path)
    return output_path


def main() -> None:
    data = load_data(DATA_PATH)
    output = generate_presentation(data, OUTPUT_PATH)
    print(f"Presentation saved to {output}")


if __name__ == "__main__":
    main()
