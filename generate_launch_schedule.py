#!/usr/bin/env python3
"""Generate RTR launch schedule presentation from CSV.

Style and layout: ai/rules/cherkizovo-presentations.md
Theme: themes/cherkizovo.py, themes/michurin.py
"""

from __future__ import annotations

import argparse
import csv
import sys
from dataclasses import dataclass, field
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

ROOT_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT_DIR))

from themes.cherkizovo import (  # noqa: E402
    COLORS,
    LOGO_LEFT_IN,
    LOGO_TOP_IN,
    LOGO_WIDTH_IN,
    SIDE_MARGIN_IN,
    SLIDE_HEIGHT_IN,
    SLIDE_WIDTH_IN,
    TABLE_LEFT_IN,
    TABLE_WIDTH_IN,
    TITLE_FONT,
    TITLE_LEFT_IN,
    TITLE_SIZE,
    TITLE_TOP_IN,
)
from themes.michurin import apply_michurin_theme  # noqa: E402

ASSETS_DIR = ROOT_DIR / "assets"
LOGO_PATH = ASSETS_DIR / "logo_0.png"

MONTH_LABELS = {
    "Jan": "01",
    "Feb": "02",
    "Mar": "03",
    "Apr": "04",
    "May": "05",
    "Jun": "06",
    "Jul": "07",
    "Aug": "08",
    "Sep": "09",
    "Oct": "10",
    "Nov": "11",
    "Dec": "12",
}

LABEL_COLUMNS = [
    ("Контур", 0.9),
    ("Блок", 0.55),
    ("Задача / направление", 2.8),
    ("Сист.", 0.45),
]

MAX_ROWS_PER_SLIDE = 9
TABLE_TOP_IN = 1.127
HEADER_ROW_HEIGHT_IN = 0.55
MIN_ROW_HEIGHT_IN = 0.42
SLIDE_TITLE_TEMPLATE = "График запуска ({page}/{total})"


def rgb(name: str) -> RGBColor:
    red, green, blue = COLORS[name]
    return RGBColor(red, green, blue)


@dataclass
class TimelineColumn:
    index: int
    year: str
    month: str

    @property
    def key(self) -> str:
        return f"{self.year}-{self.month}"

    @property
    def label(self) -> str:
        month_num = MONTH_LABELS.get(self.month, self.month[:3])
        year_short = self.year[-2:] if self.year else ""
        return f"{month_num}.{year_short}"


@dataclass
class ScheduleRow:
    kind: str  # section | subsection | task
    contour: str = ""
    block: str = ""
    task: str = ""
    system: str = ""
    marks: dict[str, str] = field(default_factory=dict)
    owner: str = ""


@dataclass
class SlideGroup:
    title: str
    rows: list[ScheduleRow]


def load_schedule(path: Path) -> tuple[list[TimelineColumn], list[ScheduleRow]]:
    with path.open(encoding="utf-8-sig") as handle:
        rows = list(csv.reader(handle))

    years = rows[1]
    months = rows[2]
    timeline: list[TimelineColumn] = []
    for index in range(8, len(months)):
        year = years[index].strip() if index < len(years) else ""
        month = months[index].strip() if index < len(months) else ""
        if year or month:
            timeline.append(TimelineColumn(index=index, year=year, month=month))

    records: list[ScheduleRow] = []
    for row in rows[3:]:
        contour = row[0].strip() if len(row) > 0 else ""
        block = row[1].strip() if len(row) > 1 else ""
        task = row[2].strip() if len(row) > 2 else ""
        system = row[6].strip() if len(row) > 6 else ""

        if not any((contour, block, task, system)):
            continue

        marks = {
            column.key: row[column.index].strip()
            for column in timeline
            if column.index < len(row) and row[column.index].strip()
        }

        if task and not contour and not block and not marks and not system:
            if task.isupper() or task.startswith("MTD") or task.startswith("RTR"):
                records.append(ScheduleRow(kind="section", task=task))
                continue
            if task in {"Себестоимость", "Бюджетный контроль", "Онлайн себестоимость"}:
                records.append(ScheduleRow(kind="subsection", task=task))
                continue
            if "Казначейство" in task or "Трансфертное" in task:
                records.append(ScheduleRow(kind="subsection", task=task))
                continue

        if contour and not task:
            records.append(ScheduleRow(kind="task", contour=contour, owner=contour))
            continue

        if not task:
            continue

        records.append(
            ScheduleRow(
                kind="task",
                contour=contour,
                block=block,
                task=task,
                system=system,
                marks=marks,
                owner=contour,
            )
        )

    return timeline, records


def group_slides(records: list[ScheduleRow]) -> list[SlideGroup]:
    groups: list[SlideGroup] = []
    current_title = "Общий график"
    current_rows: list[ScheduleRow] = []

    def flush() -> None:
        nonlocal current_rows
        if not current_rows:
            return
        for start in range(0, len(current_rows), MAX_ROWS_PER_SLIDE):
            chunk = current_rows[start : start + MAX_ROWS_PER_SLIDE]
            suffix = ""
            total_chunks = (len(current_rows) + MAX_ROWS_PER_SLIDE - 1) // MAX_ROWS_PER_SLIDE
            if total_chunks > 1:
                chunk_index = start // MAX_ROWS_PER_SLIDE + 1
                suffix = f" — {chunk_index}/{total_chunks}"
            groups.append(SlideGroup(title=f"{current_title}{suffix}", rows=chunk))
        current_rows = []

    for record in records:
        if record.kind == "section":
            flush()
            current_title = record.task
            continue
        current_rows.append(record)

    flush()
    return groups


def set_cell_border(cell, *, color: str = "000000", width: str = "6350") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    for edge in ("lnL", "lnR", "lnT", "lnB"):
        tag = qn(f"a:{edge}")
        element = tc_pr.find(tag)
        if element is None:
            element = OxmlElement(f"a:{edge}")
            tc_pr.append(element)
        element.set("w", width)
        element.set("cap", "flat")
        element.set("cmpd", "sng")
        element.set("algn", "ctr")
        fill = element.find(qn("a:solidFill"))
        if fill is None:
            fill = OxmlElement("a:solidFill")
            element.append(fill)
        srgb = fill.find(qn("a:srgbClr"))
        if srgb is None:
            srgb = OxmlElement("a:srgbClr")
            fill.append(srgb)
        srgb.set("val", color)


def style_table_borders(table) -> None:
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(cell)


def set_run_font(
    run,
    *,
    font_name: str,
    size: int,
    bold: bool = False,
    color: RGBColor | None = None,
) -> None:
    run.font.name = font_name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_cell(
    cell,
    text: str,
    *,
    font_name: str = "Calibri",
    bold: bool = False,
    size: int = 9,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    fill: RGBColor | None = None,
    font_color: RGBColor | None = None,
    valign=MSO_ANCHOR.MIDDLE,
) -> None:
    cell.text = ""
    cell.vertical_anchor = valign
    cell.margin_left = Pt(2)
    cell.margin_right = Pt(2)
    cell.margin_top = Pt(1)
    cell.margin_bottom = Pt(1)
    cell.text_frame.word_wrap = True
    paragraph = cell.text_frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_before = Pt(0)
    paragraph.space_after = Pt(0)
    paragraph.line_spacing = 1.0
    run = paragraph.add_run()
    run.text = text
    set_run_font(
        run,
        font_name=font_name,
        size=size,
        bold=bold,
        color=font_color or rgb("black"),
    )
    if fill is not None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill


def add_logo(slide) -> None:
    if LOGO_PATH.exists():
        slide.shapes.add_picture(
            str(LOGO_PATH),
            Inches(LOGO_LEFT_IN),
            Inches(LOGO_TOP_IN),
            width=Inches(LOGO_WIDTH_IN),
        )


def set_textbox(
    text_frame,
    text: str,
    *,
    font_name: str,
    size: int,
    bold: bool = False,
    color: RGBColor | None = None,
    align: PP_ALIGN = PP_ALIGN.LEFT,
) -> None:
    text_frame.clear()
    paragraph = text_frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    set_run_font(
        run,
        font_name=font_name,
        size=size,
        bold=bold,
        color=color or rgb("black"),
    )


def add_title_slide(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = rgb("dark_burgundy")

    title_box = slide.shapes.add_textbox(
        Inches(TITLE_LEFT_IN),
        Inches(2.4),
        Inches(8.5),
        Inches(1.2),
    )
    set_textbox(
        title_box.text_frame,
        "ГРАФИК ЗАПУСКА",
        font_name=TITLE_FONT,
        size=32,
        bold=True,
        color=rgb("white"),
    )

    subtitle_box = slide.shapes.add_textbox(
        Inches(TITLE_LEFT_IN),
        Inches(3.5),
        Inches(9.0),
        Inches(0.8),
    )
    set_textbox(
        subtitle_box.text_frame,
        "RTR — управленческий учёт\nMTD и продуктивный контур",
        font_name=TITLE_FONT,
        size=16,
        color=rgb("alt_row"),
    )

    if LOGO_PATH.exists():
        slide.shapes.add_picture(
            str(LOGO_PATH),
            Inches(LOGO_LEFT_IN),
            Inches(LOGO_TOP_IN),
            width=Inches(LOGO_WIDTH_IN),
        )


def estimate_row_height(record: ScheduleRow) -> float:
    if record.kind in {"section", "subsection"}:
        return 0.35
    text_len = max(len(record.task), max((len(value) for value in record.marks.values()), default=0))
    if text_len > 60:
        return 0.62
    if text_len > 30:
        return 0.5
    return MIN_ROW_HEIGHT_IN


def build_gantt_slide(
    prs: Presentation,
    *,
    group: SlideGroup,
    timeline: list[TimelineColumn],
    page_index: int,
    total_pages: int,
) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_logo(slide)

    title_box = slide.shapes.add_textbox(
        Inches(TITLE_LEFT_IN),
        Inches(TITLE_TOP_IN),
        Inches(9.5),
        Inches(0.45),
    )
    set_textbox(
        title_box.text_frame,
        f"{group.title} ({page_index}/{total_pages})",
        font_name=TITLE_FONT,
        size=TITLE_SIZE,
        bold=True,
        color=rgb("title"),
    )

    label_widths = [width for _, width in LABEL_COLUMNS]
    month_width = (TABLE_WIDTH_IN - sum(label_widths)) / max(1, len(timeline))
    column_widths = [Inches(width) for width in label_widths] + [
        Inches(month_width) for _ in timeline
    ]

    row_heights = [estimate_row_height(record) for record in group.rows]
    table_height = Inches(HEADER_ROW_HEIGHT_IN + sum(row_heights))

    table_shape = slide.shapes.add_table(
        len(group.rows) + 1,
        len(LABEL_COLUMNS) + len(timeline),
        Inches(TABLE_LEFT_IN),
        Inches(TABLE_TOP_IN),
        Inches(TABLE_WIDTH_IN),
        table_height,
    )
    table = table_shape.table

    for index, width in enumerate(column_widths):
        table.columns[index].width = width

    table.rows[0].height = Inches(HEADER_ROW_HEIGHT_IN)
    for row_index, row_height in enumerate(row_heights, start=1):
        table.rows[row_index].height = Inches(row_height)

    for column_index, (title, _) in enumerate(LABEL_COLUMNS):
        set_cell(
            table.cell(0, column_index),
            title,
            font_name="Calibri",
            bold=True,
            size=9,
            align=PP_ALIGN.CENTER,
            fill=rgb("dark_red"),
            font_color=rgb("white"),
        )

    for column_index, column in enumerate(timeline, start=len(LABEL_COLUMNS)):
        set_cell(
            table.cell(0, column_index),
            column.label,
            font_name="Calibri",
            bold=True,
            size=7,
            align=PP_ALIGN.CENTER,
            fill=rgb("dark_red"),
            font_color=rgb("white"),
        )

    for row_index, record in enumerate(group.rows, start=1):
        if record.kind == "subsection":
            for column_index in range(len(LABEL_COLUMNS) + len(timeline)):
                set_cell(
                    table.cell(row_index, column_index),
                    record.task if column_index == 2 else "",
                    font_name=TITLE_FONT,
                    bold=True,
                    size=10,
                    fill=rgb("title"),
                    font_color=rgb("white"),
                )
            continue

        fill = rgb("alt_row") if row_index % 2 == 0 else rgb("white")
        set_cell(table.cell(row_index, 0), record.contour, size=8, fill=fill)
        set_cell(table.cell(row_index, 1), record.block, size=8, fill=fill, align=PP_ALIGN.CENTER)
        set_cell(table.cell(row_index, 2), record.task, size=8, fill=fill)
        set_cell(
            table.cell(row_index, 3),
            record.system,
            size=8,
            fill=fill,
            align=PP_ALIGN.CENTER,
        )

        for column_index, column in enumerate(timeline, start=len(LABEL_COLUMNS)):
            mark = record.marks.get(column.key, "")
            if mark:
                set_cell(
                    table.cell(row_index, column_index),
                    mark,
                    size=6,
                    align=PP_ALIGN.CENTER,
                    fill=rgb("brand_red"),
                    font_color=rgb("white"),
                )
            else:
                set_cell(
                    table.cell(row_index, column_index),
                    "",
                    size=6,
                    fill=fill,
                )

    style_table_borders(table)


def build_presentation(
    timeline: list[TimelineColumn],
    records: list[ScheduleRow],
    output_path: Path,
) -> int:
    groups = group_slides(records)
    total_pages = len(groups)

    prs = Presentation()
    prs.slide_width = Inches(SLIDE_WIDTH_IN)
    prs.slide_height = Inches(SLIDE_HEIGHT_IN)
    apply_michurin_theme(prs)

    add_title_slide(prs)

    for page_index, group in enumerate(groups, start=1):
        build_gantt_slide(
            prs,
            group=group,
            timeline=timeline,
            page_index=page_index,
            total_pages=total_pages,
        )

    prs.save(output_path)
    return len(groups)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--csv",
        type=Path,
        default=Path("data/RTR-launch-schedule.csv"),
        help="Path to launch schedule CSV",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("График_запуска_RTR.pptx"),
        help="Path to generated presentation",
    )
    args = parser.parse_args()

    timeline, records = load_schedule(args.csv)
    slide_count = build_presentation(timeline, records, args.output)
    print(
        f"Created {args.output} with {len(records)} rows across "
        f"{slide_count} gantt slides (+ title slide)"
    )


if __name__ == "__main__":
    main()
